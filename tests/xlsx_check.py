# Kikan - suivi des actifs et remise de materiel signee
# Copyright (C) 2026 Gabriel Chollet
# Developpe avec l'assistance d'une IA (Claude, Anthropic).
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

"""Tests de l'import et de l'export Excel (.xlsx) : equivalence avec le CSV, fichiers hostiles, modeles, export.

Usage :
  docker exec -i -e TEST_ADMIN_USER=admin -e TEST_ADMIN_PASSWORD=... kikan-api python - < tests/xlsx_check.py

Cree des donnees ZZXL et des comptes sec_xl* (nettoyage : tests/cleanup_test_data.sql). Fichier ASCII uniquement.
Code de sortie 1 si un test echoue.
"""
import io
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
import zipfile
from datetime import date, datetime

import openpyxl

B = os.environ.get("TEST_BASE_URL", "http://localhost:8000")
ADMIN = os.environ.get("TEST_ADMIN_USER", "admin")
ADMIN_PW = os.environ.get("TEST_ADMIN_PASSWORD", "")
PW = "Sec-Test-Passw0rd-2026"
E = chr(233)
fails = 0


def show(name, ok, detail=""):
    global fails
    if not ok:
        fails += 1
    print(("PASS " if ok else "FAIL ") + name + (" | " + str(detail)[:170] if detail else ""))


def call(method, path, body=None, token=None, raw=None, ctype="application/json"):
    h = {"Content-Type": ctype}
    if token:
        h["Authorization"] = "Bearer " + token
    data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
    req = urllib.request.Request(B + path, data=data, method=method, headers=h)
    try:
        r = urllib.request.urlopen(req, timeout=120)
        return r.status, r.read(), r.headers
    except urllib.error.HTTPError as e:
        return e.code, e.read(), e.headers


def login(user, pw):
    s, b, _ = call("POST", "/auth/login", {"username": user, "password": pw})
    return json.loads(b)["access_token"] if s == 200 else None


ADM = login(ADMIN, ADMIN_PW)
if not ADM:
    sys.exit("Connexion admin impossible")


def j(method, path, body=None, token=None):
    s, b, _ = call(method, path, body, token or ADM)
    try:
        return s, json.loads(b)
    except ValueError:
        return s, b


def post_import(kind, content, query="", token=None):
    bd = "----kikanXL"
    body = (f'--{bd}\r\nContent-Disposition: form-data; name="file"; filename="x.bin"\r\n'
            "Content-Type: application/octet-stream\r\n\r\n").encode() + content + f"\r\n--{bd}--\r\n".encode()
    s, b, _ = call("POST", f"/import/{kind}{query}", token=token or ADM, raw=body, ctype="multipart/form-data; boundary=" + bd)
    try:
        return s, json.loads(b)
    except ValueError:
        return s, b


def msg(r):
    return r.get("detail", "") if isinstance(r, dict) else str(r)


def sim(kind, content):
    return post_import(kind, content, "?dry_run=true")


def wb_bytes(rows, title="Donnees", extra=None, hidden_first=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    if extra:
        for name, data in extra:
            w2 = wb.create_sheet(name)
            for r in data:
                w2.append(r)
    if hidden_first:
        ws.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rewrite_zip(raw, edit=None, add=None):
    """Copie d'un classeur avec une entree modifiee ou ajoutee."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if edit and info.filename in edit:
                data = edit[info.filename](data)
            zout.writestr(info.filename, data)
        for name, data in (add or {}).items():
            zout.writestr(name, data)
    return out.getvalue()


def ensure_account(name, role):
    s, b, _ = call("POST", "/auth/accounts", {"username": name, "password": PW, "role": role}, ADM)
    if s == 409:
        accs = json.loads(call("GET", "/auth/accounts", token=ADM)[1])
        acc = next(a for a in accs if a["username"] == name)
        call("PUT", f"/auth/accounts/{acc['id']}", {"actif": True, "role": role, "password": PW}, ADM)
        return acc["id"]
    return json.loads(b)["id"]


ids = [ensure_account("sec_xlrh", "rh"), ensure_account("sec_xllect", "lecture")]
RH, LECT = login("sec_xlrh", PW), login("sec_xllect", PW)

print("== Modeles Excel ==")
s, specs = j("GET", "/import/specs")
for sp in specs:
    s, raw, h = call("GET", f"/import/modeles/{sp['kind']}?format=xlsx", token=ADM)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb["Donnees"]
        head = [c.value for c in ws[1]]
        ok = (s == 200 and wb.sheetnames == ["Donnees", "Aide"] and head == [c["nom"] for c in sp["colonnes"]]
              and ws.freeze_panes == "A2" and ws.max_row >= 1 + len(sp["exemples"])
              and "spreadsheetml" in h.get("Content-Type", "") and "attachment" in h.get("Content-Disposition", ""))
    except Exception as exc:  # noqa: BLE001
        ok = False
        head = str(exc)
    show(f"modele {sp['kind']} : classeur valide (feuilles Donnees et Aide, en-tetes, exemples, volets figes)", ok, head if not ok else "")
    aide = wb["Aide"]
    show(f"modele {sp['kind']} : feuille Aide decrit chaque colonne", [aide.cell(row=r, column=1).value for r in range(2, 2 + len(sp["colonnes"]))] == [c["nom"] for c in sp["colonnes"]])
s, raw, _ = call("GET", "/import/modeles/utilisateurs?format=xlsx", token=ADM)
wb = openpyxl.load_workbook(io.BytesIO(raw))
ws = wb["Donnees"]
dvs = ws.data_validations.dataValidation
show("modele : listes deroulantes (statut)", any("prevu" in (d.formula1 or "") for d in dvs))
show("modele : colonnes au format texte (zeros initiaux preserves)", ws["A5"].number_format == "@" and ws["B200"].number_format == "@")
for sp in specs:
    s, raw, _ = call("GET", f"/import/modeles/{sp['kind']}?format=xlsx", token=ADM)
    s, r = sim(sp["kind"], raw)
    show(f"modele {sp['kind']} : reimportable (lu comme un classeur, colonnes reconnues)", s == 200 and r["total"] == len(sp["exemples"]) and r["colonnes_ignorees"] == [] and r["encodage"] == "Excel (.xlsx)", r if s != 200 else "")
show("modele : format invalide -> 422", call("GET", "/import/modeles/entites?format=ods", token=ADM)[0] == 422)
show("modele CSV toujours disponible", call("GET", "/import/modeles/entites", token=ADM)[0] == 200)
for name, tok in (("RH", RH), ("lecture", LECT)):
    show(f"{name} : modele Excel refuse (403)", call("GET", "/import/modeles/entites?format=xlsx", token=tok)[0] == 403)
    show(f"{name} : import Excel refuse (403)", post_import("entites", wb_bytes([["code", "nom"], ["ZZXL-R", "x"]]), "?dry_run=true", tok)[0] == 403)
show("sans jeton : refuse (401)", call("GET", "/import/modeles/entites?format=xlsx")[0] == 401)

print("== Equivalence CSV et Excel ==")
csv_ent = b"code;nom;code_parent;actif\nZZXL-A;ZZXL Groupe;;oui\nZZXL-B;ZZXL Filiale;ZZXL-A;non\n;sans code;;oui\n"
x_ent = wb_bytes([["code", "nom", "code_parent", "actif"], ["ZZXL-A", "ZZXL Groupe", None, "oui"], ["ZZXL-B", "ZZXL Filiale", "ZZXL-A", False], [None, "sans code", None, True]])
s1, rc = sim("entites", csv_ent)
s2, rx = sim("entites", x_ent)
show("entites : meme rapport en CSV et en Excel", s1 == s2 == 200 and (rc["imported"], rc["skipped"], rc["total"], rc["errors"]) == (rx["imported"], rx["skipped"], rx["total"], rx["errors"]), (rc["errors"], rx["errors"]))
show("Excel : feuille et format renvoyes", rx["encodage"] == "Excel (.xlsx)" and rx["feuille"] == "Donnees" and rx["feuilles_ignorees"] == 0)
show("Excel : empreinte du fichier", len(rx["sha256"]) == 64)
csv_usr = ("matricule;nom;prenom;service;date_arrivee;statut\n"
           f"ZZXL007;Dur{E}and;Cl{E}mence;Compta;2026-01-05;actif\n"
           "ZZXL1001;Martin;Paul;Logistique;2026-02-10;prevu\n").encode("cp1252")
x_usr = wb_bytes([["matricule", "nom", "prenom", "service", "date_arrivee", "statut"],
                  ["ZZXL007", f"Dur{E}and", f"Cl{E}mence", "Compta", date(2026, 1, 5), "actif"],
                  ["ZZXL1001", "Martin", "Paul", "Logistique", datetime(2026, 2, 10), "prevu"]])
s1, rc = sim("utilisateurs", csv_usr)
s2, rx = sim("utilisateurs", x_usr)
show("salaries : accents et dates Excel (cellules date) acceptes, comme le CSV", s2 == 200 and rx["imported"] == rc["imported"] == 2 and rx["errors"] == [], rx["errors"])
post_import("utilisateurs", x_usr)
s, lst = j("GET", "/users?q=ZZXL")
by = {u["matricule"]: u for u in lst}
show("import reel Excel : valeurs lues correctement", by["ZZXL007"]["nom"] == f"Dur{E}and" and by["ZZXL007"]["date_arrivee"] == "2026-01-05"
     and by["ZZXL1001"]["statut"] == "prevu" and by["ZZXL1001"]["date_arrivee"] == "2026-02-10", {k: (v["nom"], v["date_arrivee"]) for k, v in by.items()})
x_num = wb_bytes([["matricule", "nom", "prenom"], [4242, "Num", "ZZXL-Eric"], [7.0, "Flottant", "ZZXL-Zoe"], [3.5, "Decimal", "ZZXL-Dan"], ["007", "Texte", "ZZXL-Tom"]])
s, r = post_import("utilisateurs", x_num)
s, lst = j("GET", "/users?q=Num")
mats = {u["nom"]: u["matricule"] for u in j("GET", "/users")[1] if u["nom"] in ("Num", "Flottant", "Decimal", "Texte")}
show("nombres : entier sans .0, flottant entier, decimal, zeros initiaux du texte", mats == {"Num": "4242", "Flottant": "7", "Decimal": "3.5", "Texte": "007"}, mats)
x_bool = wb_bytes([["nom", "application", "actif"], ["x", "x", True]])
x_apps = wb_bytes([["nom", "description", "actif"], ["ZZXL-App", "d", True], ["ZZXL-App2", "d", False]])
post_import("applications", x_apps)
s, apps = j("GET", "/applications")
act = {a["nom"]: a["actif"] for a in apps if a["nom"].startswith("ZZXL")}
show("booleens Excel lus comme oui/non", act == {"ZZXL-App": True, "ZZXL-App2": False}, act)

print("== Numerotation, feuilles, formules ==")
x_gap = wb_bytes([["code", "nom"], ["ZZXL-G1", "ok"], [None, None], [None, "sans code"], ["ZZXL-G2", "ok2"]])
s, r = sim("entites", x_gap)
show("lignes vides ignorees, numeros de ligne = lignes du classeur", s == 200 and r["total"] == 3 and [e["ligne"] for e in r["errors"]] == [4], r["errors"])
show("detail : numeros de ligne du classeur", [l["ligne"] for l in r["lignes"]] == [2, 5], r["lignes"])
x_multi = wb_bytes([["code", "nom"], ["ZZXL-M1", "premiere feuille"]], extra=[("Autre", [["a", "b"], [1, 2]])])
s, r = sim("entites", x_multi)
show("plusieurs feuilles : la premiere est lue, les autres signalees", r["feuille"] == "Donnees" and r["feuilles_ignorees"] == 1 and r["imported"] == 1, r.get("feuille"))
x_hid = wb_bytes([["junk", "junk"], ["a", "b"]], title="Cachee", extra=[("Visible", [["code", "nom"], ["ZZXL-V1", "visible"]])], hidden_first=True)
s, r = sim("entites", x_hid)
show("feuille cachee ignoree : premiere feuille visible lue", s == 200 and r["feuille"] == "Visible" and r["imported"] == 1, msg(r) if s != 200 else r.get("feuille"))
x_f = wb_bytes([["code", "nom"], ["ZZXL-F1", "valeur"], ["ZZXL-F2", "=1+1"], ["ZZXL-F3", "autre"]])
s, r = sim("entites", x_f)
show("formule sans valeur enregistree : ligne en erreur avec consigne", s == 200 and [e["ligne"] for e in r["errors"]] == [3] and "enregistrez-le" in r["errors"][0]["message"], r["errors"])
show("... les autres lignes restent valides", r["imported"] == 2)
s, r = post_import("entites", x_f, "?strict=true")
show("... import reel strict refuse (409), rien cree", s == 409 and "ZZXL-F1" not in {e["code"] for e in j("GET", "/entites")[1]})


def with_cached(data):
    return re.sub(rb"<f>1\+1</f><v\s*/>|<f>1\+1</f><v></v>", b"<f>1+1</f><v>2</v>", data)


x_cached = rewrite_zip(x_f, edit={"xl/worksheets/sheet1.xml": with_cached})
s, r = sim("entites", x_cached)
show("formule AVEC valeur enregistree par Excel : valeur lue, jamais recalculee", s == 200 and r["errors"] == [] and r["imported"] == 3 and any(l["message"].endswith("2") for l in r["lignes"]), r["lignes"])
wbx = openpyxl.Workbook()
wbx.active.append(["code", "nom"])
cell_ws = wbx.active
cell_ws.append(["ZZXL-T1", "x"])
cell_ws.cell(row=2, column=2, value="=cmd|' /C calc'!A0")
cell_ws.cell(row=2, column=2).data_type = "s"
buf = io.BytesIO()
wbx.save(buf)
post_import("entites", buf.getvalue())
s, ents = j("GET", "/entites")
show("texte commencant par = : importe tel quel (texte, jamais evalue)", any(e["nom"].startswith("=cmd") for e in ents if e["code"] == "ZZXL-T1"))
show("equivalence : colonne obligatoire absente -> meme message qu'en CSV",
     post_import("entites", wb_bytes([["code", "libelle"], ["a", "b"]]))[0] == 400 and "nom" in msg(post_import("entites", wb_bytes([["code", "libelle"], ["a", "b"]]))[1]))

print("== Fichiers hostiles ==")
def refused(name, content, codes=(400,), needle=None):
    s, r = sim("entites", content)
    ok = s in codes and (needle is None or needle in msg(r))
    show(name, ok, f"HTTP {s} {msg(r)}")

good = wb_bytes([["code", "nom"], ["ZZXL-H", "h"]])
refused("ancien format .xls (OLE) refuse, message clair", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600, needle=".xls")
refused("classeur chiffre (OLE) refuse", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + os.urandom(2000), needle="mot de passe")
refused("macros (vbaProject.bin) refusees", rewrite_zip(good, add={"xl/vbaProject.bin": b"binaire"}), needle="macro")
refused("liens externes refuses", rewrite_zip(good, add={"xl/externalLinks/externalLink1.xml": b"<x/>"}), needle="liens")
refused("archive ZIP qui n'est pas un classeur", (lambda b: (zipfile.ZipFile(b, "w").writestr("a.txt", "x"), b.getvalue())[1])(io.BytesIO()), needle="n'est pas un classeur")
refused("classeur tronque", good[:400])
refused("faux classeur (en-tete ZIP seul)", b"PK\x03\x04" + b"\x00" * 100)
bomb = rewrite_zip(good, add={"xl/media/bombe.bin": b"\x00" * (200 * 1024 * 1024)})
t0 = time.time()
refused("bombe de decompression (200 Mo decompresses)", bomb, needle="volumineuse")
show("   ... refusee rapidement, sans lecture du classeur", time.time() - t0 < 10, f"{time.time() - t0:.1f} s")
lol = (b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">'
       b'<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">]><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&lol3;</t></is></c></row></sheetData></worksheet>')
t0 = time.time()
refused("bombe XML (entites imbriquees) refusee", rewrite_zip(good, edit={"xl/worksheets/sheet1.xml": lambda d: lol}), needle="illisible")
show("   ... refusee rapidement", time.time() - t0 < 10, f"{time.time() - t0:.1f} s")
refused("feuille vide", wb_bytes([]), needle="vide")
refused("en-tetes seuls, aucune ligne de donnees", wb_bytes([["code", "nom"]]), needle="aucune ligne")
refused("premiere ligne sans en-tetes", wb_bytes([[None, None], ["a", "b"]]), needle="en-tetes")
big = wb_bytes([["code", "nom"]] + [[f"ZZXL-N{i}", "n"] for i in range(20001)])
s, r = sim("entites", big)
show("plus de 20000 lignes -> 413", s == 413, f"HTTP {s}")
padded = rewrite_zip(good, add={"xl/media/gros.bin": os.urandom(6 * 1024 * 1024)})
s, r = sim("entites", padded)
show("fichier de plus de 5 Mo -> 413", s == 413, f"HTTP {s}")
refused("fichier binaire quelconque", os.urandom(300))

print("== Import reel, tout ou rien, empreinte ==")
x_ok = wb_bytes([["code", "nom"], ["ZZXL-OK1", "un"], ["ZZXL-OK2", "deux"]])
s, r = sim("entites", x_ok)
s2, bad = post_import("entites", x_ok, "?attendu=" + "0" * 64)
show("fichier different de la simulation -> 409", s2 == 409)
s3, real = post_import("entites", x_ok, "?strict=true&attendu=" + r["sha256"])
show("import reel Excel verifie par empreinte", s3 == 200 and real["imported"] == 2 and {"ZZXL-OK1", "ZZXL-OK2"} <= {e["code"] for e in j("GET", "/entites")[1]}, real if s3 != 200 else "")
s, again = sim("entites", x_ok)
show("reimport : tout ignore (idempotent)", again["imported"] == 0 and again["skipped"] == 2)
x_mat = wb_bytes([["identifiant_interne", "type", "marque", "modele", "numero_serie", "statut"],
                  ["ZZXL-PC1", "PC portable", "zzxl dell", "Latitude", "ZZXL-SN1", "stock"],
                  ["ZZXL-PC2", "PC portable", "ZZXL Dell", "latitude", "ZZXL-SN2", "stock"]])
s, r = post_import("materiel", x_mat)
s, allr = j("GET", "/assets?statut=stock")
mine = [a for a in allr if (a.get("numero_serie") or "").startswith("ZZXL-SN")]
show("materiel Excel : memes regles qu'en CSV (graphie canonique des marques)", r["imported"] == 2 and len({a["marque"] for a in mine}) == 1, {a["marque"] for a in mine})

print("== Export Excel du rapport ==")
s, usr = j("POST", "/users", {"prenom": "Zoe", "nom": "=cmd|calc", "matricule": "ZZXL-U9", "statut": "actif"})
s, ast = j("POST", "/assets", {"type": "PC", "marque": "ZZXL", "modele": "M", "numero_serie": "ZZXL-SN9", "statut": "stock"})
j("POST", "/assignments", {"user_id": usr["id"], "asset_ids": [ast["id"]], "sans_charte": True})
s, raw, h = call("GET", "/reports/holdings?format=xlsx", token=ADM)
wb = openpyxl.load_workbook(io.BytesIO(raw))
ws = wb["Detentions"]
s, js = j("GET", "/reports/holdings?format=json")
head = [c.value for c in ws[1]]
show("export : classeur valide, en-tetes du rapport", s == 200 and head[:3] == ["Entite", "Service", "Matricule"] and head[-1] == "Remise" and "spreadsheetml" in h.get("Content-Type", ""), head)
show("export : une ligne par detention, comme le JSON", ws.max_row - 1 == js["count"], (ws.max_row - 1, js["count"]))
show("export : filtre automatique et volets figes", ws.auto_filter.ref is not None and ws.freeze_panes == "A2")
mine = [r for r in ws.iter_rows(min_row=2) if r[2].value == "ZZXL-U9"]
show("export : contenu '=cmd|calc' stocke comme TEXTE (jamais une formule)", len(mine) == 1 and mine[0][3].value == "=cmd|calc" and mine[0][3].data_type == "s", mine[0][3].data_type if mine else None)
show("export : date de remise typee (vraie date Excel)", isinstance(mine[0][11].value, (datetime, date)), type(mine[0][11].value).__name__)
for name, tok in (("RH", RH), ("lecture", LECT)):
    s_csv = call("GET", "/reports/holdings?format=csv", token=tok)[0]
    s_x = call("GET", "/reports/holdings?format=xlsx", token=tok)[0]
    show(f"export : memes droits que le CSV pour {name}", s_x == s_csv == 200, (s_csv, s_x))
show("export : sans jeton refuse (401)", call("GET", "/reports/holdings?format=xlsx")[0] == 401)
show("export : format inconnu -> 422", call("GET", "/reports/holdings?format=ods", token=ADM)[0] == 422)
ent_id = j("GET", "/entites")[1][0]["id"]
s, raw2, _ = call("GET", f"/reports/holdings?format=xlsx&entite_id={ent_id}", token=ADM)
s, js2 = j("GET", f"/reports/holdings?format=json&entite_id={ent_id}")
show("export : filtres (entite) appliques comme pour le CSV", openpyxl.load_workbook(io.BytesIO(raw2))["Detentions"].max_row - 1 == js2["count"])

for aid in ids:
    call("PUT", f"/auth/accounts/{aid}", {"actif": False}, ADM)
print(f"\n{'ECHEC' if fails else 'OK'} : {fails} test(s) en echec")
sys.exit(1 if fails else 0)
